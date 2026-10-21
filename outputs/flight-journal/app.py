"""Local flight journal. Read-only Excel source; no cloud services required."""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import mimetypes
import os
import re
import threading
from datetime import date, datetime, time, timezone
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlsplit
from zoneinfo import ZoneInfo

import openpyxl
from openpyxl.utils.datetime import from_excel

BASE = Path(__file__).resolve().parent
WEB = BASE / "web"
CONFIG = BASE / "config.json"
AIRPORTS = json.loads((WEB / "data" / "airports.json").read_text(encoding="utf-8"))
AIRPORTS.update(json.loads((WEB / "data" / "historical-airports.json").read_text(encoding="utf-8")))
HEADERS = ["起飞日期", "航班号", "出发机场", "到达机场", "起飞时间", "到达日期", "到达时间", "航空公司", "订单/预订号", "凭证状态", "乘坐/导入判定", "备注", "来源邮箱", "邮件日期", "原邮件链接", "补充来源链接"]
FIELDS = ["date", "number", "originLabel", "destinationLabel", "departure", "arrivalDate", "arrival", "airline", "booking", "evidence", "status", "notes", "mailbox", "mailDate", "source", "extraSources"]
FLOWN = {"已乘坐", "推定已乘坐", "已导入"}


def clean(value):
    return "" if value is None else str(value).strip()


def parse_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (float, int)) and value > 1:
        return from_excel(value).date().isoformat()
    text = clean(value)
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text[:10], fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError("无法识别日期")


def parse_time(value):
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, (float, int)) and 0 <= value < 1:
        minutes = round(value * 1440) % 1440
        return f"{minutes // 60:02}:{minutes % 60:02}"
    match = re.fullmatch(r"(\d{1,2}):(\d{2})(?::\d{2})?", clean(value))
    if match and int(match[1]) < 24 and int(match[2]) < 60:
        return f"{int(match[1]):02}:{match[2]}"
    return None


def airport_code(label):
    match = re.match(r"^([A-Z]{3})(?:\b|\s)", clean(label).upper())
    return match[1] if match else ""


def distance_km(origin, destination):
    if not origin or not destination:
        return None
    lat1, lat2 = math.radians(origin["lat"]), math.radians(destination["lat"])
    dlat = lat2 - lat1
    dlon = math.radians(destination["lon"] - origin["lon"])
    h = math.sin(dlat / 2) ** 2 + math.cos(lat1) * math.cos(lat2) * math.sin(dlon / 2) ** 2
    return round(6371.0088 * 2 * math.asin(math.sqrt(min(1, h))))


def duration_minutes(row, origin, destination):
    if not all((origin, destination, row["date"], row["arrivalDate"], row["departure"], row["arrival"])):
        return None
    try:
        departure = datetime.fromisoformat(f'{row["date"]}T{row["departure"]}').replace(tzinfo=ZoneInfo(origin["tz"]))
        arrival = datetime.fromisoformat(f'{row["arrivalDate"]}T{row["arrival"]}').replace(tzinfo=ZoneInfo(destination["tz"]))
        minutes = round((arrival.astimezone(timezone.utc) - departure.astimezone(timezone.utc)).total_seconds() / 60)
        return minutes if 0 < minutes <= 1440 else None
    except (ValueError, KeyError):
        return None


def safe_links(value):
    return [url.strip() for url in clean(value).splitlines() if urlsplit(url.strip()).scheme in {"https", "http"}]


def read_flights(sheet, warnings):
    rows = iter(sheet.iter_rows(values_only=True))
    columns = None
    header_row = 0
    for index, values in enumerate(rows, 1):
        labels = [clean(v) for v in values]
        if all(h in labels for h in ("起飞日期", "出发机场", "到达机场")):
            columns = {field: labels.index(header) if header in labels else None for field, header in zip(FIELDS, HEADERS)}
            header_row = index
            break
        if index >= 20:
            break
    if columns is None:
        raise ValueError(f'「{sheet.title}」未找到起飞日期、出发机场和到达机场列')
    flights = []
    for index, values in enumerate(rows, header_row + 1):
        if not any(v is not None for v in values):
            continue
        item = {k: values[i] if i is not None and i < len(values) else None for k, i in columns.items()}
        if not any(item[k] for k in ("date", "number", "originLabel", "destinationLabel")):
            continue
        try:
            item["date"] = parse_date(item["date"])
        except (ValueError, OverflowError):
            item["date"] = None
        if not item["date"]:
            warnings.append(f"{sheet.title} 第 {index} 行缺少有效起飞日期，暂未计入。")
            continue
        for field in ("arrivalDate", "mailDate"):
            try:
                item[field] = parse_date(item[field])
            except (ValueError, OverflowError):
                item[field] = None
        for field in ("departure", "arrival"):
            item[field] = parse_time(item[field])
        for field in FIELDS:
            if field not in {"date", "arrivalDate", "mailDate", "departure", "arrival"}:
                item[field] = clean(item[field])
        item["origin"] = airport_code(item["originLabel"])
        item["destination"] = airport_code(item["destinationLabel"])
        item["originCity"] = item["originLabel"][3:].strip() or item["origin"]
        item["destinationCity"] = item["destinationLabel"][3:].strip() or item["destination"]
        source_airport, dest_airport = AIRPORTS.get(item["origin"]), AIRPORTS.get(item["destination"])
        item["distance"] = distance_km(source_airport, dest_airport)
        item["duration"] = duration_minutes(item, source_airport, dest_airport)
        item["inferred"] = "推定" in item["evidence"] or item["status"] == "推定已乘坐"
        item["flown"] = sheet.title == "航班记录" and item["status"] in FLOWN
        item["excluded"] = sheet.title == "取消及旧票" or item["status"] in {"不导入", "未乘坐", "暂不导入"}
        item["links"] = list(dict.fromkeys(safe_links(item["source"]) + safe_links(item["extraSources"])))
        item["sheet"] = sheet.title
        item["row"] = index
        identity = "|".join([sheet.title, item["date"], item["number"], item["origin"], item["destination"], item["booking"], str(index)])
        item["id"] = hashlib.sha256(identity.encode()).hexdigest()[:16]
        flights.append(item)
    return flights


def parse_workbook(content):
    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    warnings = []
    try:
        if "航班记录" not in workbook.sheetnames:
            raise ValueError("Excel 缺少「航班记录」工作表")
        flights = read_flights(workbook["航班记录"], warnings)
        canceled = read_flights(workbook["取消及旧票"], warnings) if "取消及旧票" in workbook.sheetnames else []
        clues = []
        if "待补充线索" in workbook.sheetnames:
            for index, row in enumerate(workbook["待补充线索"].iter_rows(min_row=5, values_only=True), 5):
                if len(row) > 4 and row[2]:
                    clues.append({"title": clean(row[2]), "detail": clean(row[3]), "context": clean(row[4]), "row": index})
        codes = {r[k] for r in flights + canceled for k in ("origin", "destination")}
        missing = sorted(code for code in codes if code and code not in AIRPORTS)
        if missing:
            warnings.append("以下机场暂无坐标，不计入里程与地图：" + "、".join(missing))
        return {"flights": flights, "canceled": canceled, "clues": clues, "airports": {c: AIRPORTS[c] for c in codes if c in AIRPORTS}, "warnings": warnings}
    finally:
        workbook.close()


class JournalStore:
    def __init__(self, workbook_path=None):
        self.override = Path(workbook_path).resolve() if workbook_path else None
        self.lock = threading.RLock()
        self.signature = None
        self.last_good = None
        self.last_path = None

    def source_path(self):
        if self.override:
            return self.override
        config = json.loads(CONFIG.read_text(encoding="utf-8-sig"))
        return (BASE / config["workbook"]).resolve()

    def snapshot(self, force=False):
        with self.lock:
            now = datetime.now(timezone.utc).isoformat()
            try:
                source = self.source_path()
                self.last_path = source
                stat = source.stat()
                signature = (str(source), stat.st_mtime_ns, stat.st_size)
                if force or signature != self.signature or self.last_good is None:
                    content = source.read_bytes()
                    after = source.stat()
                    if (stat.st_mtime_ns, stat.st_size) != (after.st_mtime_ns, after.st_size):
                        raise OSError("Excel 正在保存，请稍候")
                    data = parse_workbook(content)
                    data["revision"] = hashlib.sha256(content).hexdigest()[:20]
                    data["source"] = {"name": source.name, "path": str(source), "modifiedAt": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat()}
                    self.last_good, self.signature = data, signature
                result = dict(self.last_good)
                result["sync"] = {"state": "ok", "checkedAt": now, "message": "Excel 已同步", "pollSeconds": 3}
                return result
            except Exception as error:
                self.signature = None  # Retry on every poll; never replace good data with a partial read.
                result = dict(self.last_good) if self.last_good else {"flights": [], "canceled": [], "clues": [], "airports": {}, "warnings": [], "revision": "unavailable", "source": {"name": self.last_path.name if self.last_path else "Excel", "path": str(self.last_path or ""), "modifiedAt": None}}
                if isinstance(error, FileNotFoundError):
                    message = "找不到 Excel 文件，请检查数据源路径；恢复文件后会自动重试。"
                elif isinstance(error, PermissionError):
                    message = "Excel 暂时无法读取，可能正在保存或被独占打开；稍后自动重试。"
                elif isinstance(error, ValueError):
                    message = str(error)
                else:
                    message = "Excel 暂未完整保存或格式无法读取，稍后自动重试。"
                result["sync"] = {"state": "stale" if self.last_good else "error", "checkedAt": now, "message": message, "pollSeconds": 3}
                return result


class JournalHandler(BaseHTTPRequestHandler):
    store: JournalStore

    def log_message(self, fmt, *args):
        if not args or str(args[1] if len(args) > 1 else "") not in {"200", "304"}:
            super().log_message(fmt, *args)

    def trusted_request(self):
        host = self.headers.get("Host", "").split(":")[0]
        origin = self.headers.get("Origin")
        return host in {"127.0.0.1", "localhost"} and (not origin or urlsplit(origin).hostname in {"127.0.0.1", "localhost"})

    def respond(self, body, content_type="application/json; charset=utf-8", status=200):
        if not isinstance(body, bytes):
            body = json.dumps(body, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Referrer-Policy", "no-referrer")
        self.send_header("Content-Security-Policy", "default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' data:; connect-src 'self'; frame-ancestors 'none'; base-uri 'none'")
        self.end_headers()
        try:
            self.wfile.write(body)
        except (BrokenPipeError, ConnectionResetError, ConnectionAbortedError):
            pass

    def do_GET(self):
        if not self.trusted_request():
            return self.respond({"error": "Local access only"}, status=403)
        parsed = urlsplit(self.path)
        if parsed.path == "/api/health":
            return self.respond({"app": "flight-journal", "version": 1})
        if parsed.path == "/api/flights":
            return self.respond(self.store.snapshot(force="force=1" in parsed.query))
        path = "index.html" if parsed.path == "/" else parsed.path.lstrip("/")
        file = (WEB / path).resolve()
        if not file.is_relative_to(WEB) or not file.is_file():
            return self.respond({"error": "Not found"}, status=404)
        mime = {".js": "text/javascript; charset=utf-8", ".css": "text/css; charset=utf-8", ".html": "text/html; charset=utf-8", ".json": "application/json; charset=utf-8"}.get(file.suffix, mimetypes.guess_type(str(file))[0] or "application/octet-stream")
        return self.respond(file.read_bytes(), mime)

    def do_POST(self):
        if not self.trusted_request() or self.headers.get("X-Flight-Journal") != "1":
            return self.respond({"error": "Local access only"}, status=403)
        if self.path == "/api/open-workbook":
            try:
                source = self.store.source_path()
                if not source.is_file() or source.suffix.lower() != ".xlsx":
                    raise ValueError("找不到 Excel 文件")
                os.startfile(str(source))
                return self.respond({"ok": True})
            except Exception:
                return self.respond({"error": "无法自动打开 Excel，请从数据源页面复制文件路径后打开。"}, status=400)
        return self.respond({"error": "Not found"}, status=404)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--port", type=int, default=8765)
    parser.add_argument("--workbook", type=Path)
    args = parser.parse_args()
    JournalHandler.store = JournalStore(args.workbook)
    server = ThreadingHTTPServer(("127.0.0.1", args.port), JournalHandler)
    server.daemon_threads = True
    print(f"Flight Journal ready: http://127.0.0.1:{args.port}", flush=True)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
