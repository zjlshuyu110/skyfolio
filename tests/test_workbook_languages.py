"""Workbook formats must give the same totals in either language."""
import importlib.util
import io
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('journal', ROOT / 'outputs/flight-journal/app.py')
app = importlib.util.module_from_spec(spec)
spec.loader.exec_module(app)


def content(english=True):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'Flights' if english else '航班记录'
    headers = ['Status', 'Destination', 'Departure date', 'Origin', 'Flight number', 'Evidence'] if english else ['乘坐/导入判定', '到达机场', '起飞日期', '出发机场', '航班号', '凭证状态']
    sheet.append(headers)
    statuses = ['fLoWn', 'Inferred flown', 'Imported', 'Not flown', 'On hold', 'Canceled'] if english else ['已乘坐', '推定已乘坐', '已导入', '未乘坐', '暂不导入', '不导入']
    for status in statuses:
        sheet.append([status, 'JFK', '2025-01-10', 'SFO', 'DJ100', ''])
    sheet.append(['Flown', 'JFK', 'bad date', 'SFO', 'DJ101', ''])
    canceled = book.create_sheet('Canceled' if english else '取消及旧票')
    canceled.append(headers)
    canceled.append([statuses[0], 'JFK', '2025-01-10', 'SFO', 'DJ102', ''])
    return book


def parse(book):
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return app.parse_workbook(stream.getvalue())


class WorkbookLanguageTests(unittest.TestCase):
    def test_equivalent_totals_with_reordered_headers(self):
        for english in (True, False):
            with self.subTest(english=english):
                data = parse(content(english))
                self.assertEqual(len(data['flights']), 6)
                self.assertEqual(sum(f['flown'] for f in data['flights']), 3)
                self.assertEqual(sum(f['excluded'] for f in data['flights']), 3)
                self.assertTrue(data['flights'][1]['inferred'])
                self.assertFalse(data['canceled'][0]['flown'])
                self.assertTrue(data['canceled'][0]['excluded'])
                self.assertEqual(len(data['warnings']), 1)

    def test_rejects_duplicate_language_sheets(self):
        book = content()
        book.create_sheet('航班记录')
        with self.assertRaisesRegex(ValueError, 'Duplicate sheets'):
            parse(book)

    def test_rejects_duplicate_language_columns(self):
        book = content()
        book.active.cell(1, 7, '起飞日期')
        with self.assertRaisesRegex(ValueError, 'Duplicate columns'):
            parse(book)

    def test_english_readme_has_no_chinese_characters(self):
        self.assertNotRegex((ROOT / 'README.md').read_text(encoding='utf-8'), r'[\u3400-\u9fff]')
